import os
import re
import io
import streamlit as st
from docx import Document
from yandex_cloud_ml_sdk import YCloudML

# Включаем всё нужное и загружаем настройки

st.set_page_config(page_title="Вектор 2035", layout="wide")

# Ищем файлик со стилями чтобы приложение было красивым
if os.path.exists("style.css"):
    with open("style.css", encoding="utf-8") as f:
        st.markdown(f"<style>{f.read()}</style>", unsafe_allow_html=True)

# Берем секретные ключики от Яндекса они спрятаны в файлике секретов
try:
    folder_id = st.secrets["YANDEX_FOLDER_ID"]
    api_key = st.secrets["YANDEX_API_KEY"]
except KeyError:
    st.error("Ошибка: не найдены ключи YANDEX_FOLDER_ID или YANDEX_API_KEY в secrets.toml")
    st.stop()

# Подключаем умную нейросеть

@st.cache_resource
def get_model():
    sdk = YCloudML(folder_id=folder_id, auth=api_key)
    return sdk.models.completions("yandexgpt").configure(temperature=0.6, max_tokens=2000)

def ask_ai(prompt):
    # Отправляем вопросик роботу и ждем ответ
    try:
        for ans in get_model().run(prompt):
            return ans.text.strip()
    except Exception as e:
        return f"Ошибка: {e}"

# Делаем память чтобы ничего не потерялось

if "app_data" not in st.session_state:
    st.session_state.app_data = {
        "theme": "", "category": "Инновации / стартап", "description": "",
        "tech": "", "audience": "", "ideas": [], "selected": "", 
        "plan": "", "risks": ""
    }
d = st.session_state.app_data

if "theme_error" not in st.session_state:
    st.session_state.theme_error = False

# Маленькие команды которые помогают переключать странички
def go_to_tab(tab_name):
    st.session_state.active_tab = tab_name

def reset_project():
    # Стираем всё чтобы начать с чистого листа
    st.session_state.app_data = {
        "theme": "", "category": "Инновации / стартап", "description": "",
        "tech": "", "audience": "", "ideas": [], "selected": "", 
        "plan": "", "risks": ""
    }
    st.session_state.active_tab = "2. Отправка темы"

def save_theme():
    # Проверяем не забыл ли пользователь написать тему
    if not st.session_state.app_data["theme"].strip():
        st.session_state.theme_error = True
    else:
        st.session_state.theme_error = False
        st.session_state.gen_ideas = True
        st.session_state.gen_plan = False
        st.session_state.gen_risks = False
        st.session_state.active_tab = "3. Идеи"

def select_idea(idea_text):
    # Запоминаем выбранную идею и запускаем робота писать план
    st.session_state.app_data["selected"] = idea_text
    st.session_state.gen_plan = True
    st.session_state.gen_risks = True
    st.session_state.active_tab = "4. План подготовки"

# Функция чтобы скачать красивый файлик ворд
def generate_docx(data):
    doc = Document()
    doc.add_heading('Проект "Вектор 2035"', 0)
    
    doc.add_heading('Параметры конкурса', level=1)
    doc.add_paragraph(f"Тема: {data['theme']}")
    doc.add_paragraph(f"Тип: {data['category']}")
    if data['description']: doc.add_paragraph(f"Описание: {data['description']}")
    if data['tech']: doc.add_paragraph(f"Технологии/материалы: {data['tech']}")
    if data['audience']: doc.add_paragraph(f"Аудитория: {data['audience']}")
    
    doc.add_heading('Сгенерированные идеи', level=1)
    if data['ideas']:
        for i, idea in enumerate(data['ideas'], 1):
            doc.add_paragraph(f"{i}. {idea}")
    else:
        doc.add_paragraph("Идеи отсутствуют.")
        
    doc.add_heading('Выбранная идея', level=1)
    doc.add_paragraph(data['selected'] if data['selected'] else "Идея не выбрана.")
    
    doc.add_heading('План подготовки', level=1)
    doc.add_paragraph(data['plan'] if data['plan'] else "План не сгенерирован.")
    
    doc.add_heading('Риски реализации', level=1)
    doc.add_paragraph(data['risks'] if data['risks'] else "Риски не сгенерированы.")
    
    # Складываем всё в виртуальную коробочку чтобы отдать браузеру
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# Кнопочки сверху для переключения страничек

TABS = ["1. План проекта", "2. Отправка темы", "3. Идеи", "4. План подготовки", "5. Риски"]
if "active_tab" not in st.session_state:
    st.session_state.active_tab = TABS[0]

active = st.radio("Меню", TABS, key="active_tab", horizontal=True, label_visibility="collapsed")

# Тут робот думает и придумывает пока мы ждем

if st.session_state.get("gen_ideas"):
    st.session_state.gen_ideas = False
    with st.spinner("Включаю креативность, придумываю идеи..."):
        # Собираем в кучу всё что написал пользователь
        context = f"Тема конкурса: {d['theme']}\nТип конкурса: {d['category']}\n"
        if d['description']: context += f"Описание/требования: {d['description']}\n"
        if d['tech']: context += f"Желаемые технологии/материалы: {d['tech']}\n"
        
        prompt = f"""
        Ты — эксперт, помогающий участнику подготовиться к конкурсу. 
        {context}
        Твоя задача — предложить ровно 5 идей для УЧАСТИЯ в этом конкурсе, а не для его организации. 
        Это могут быть проекты, поделки, рисунки, исследования или другие работы, которые участник может представить жюри.
        Учитывай тип конкурса и указанные требования. 
        Ответ верни в виде нумерованного списка с названиями и кратким описанием.
        """
        raw = ask_ai(prompt)
        d["ideas"] = [line for line in raw.split("\n") if line.strip()]

if st.session_state.get("gen_plan"):
    st.session_state.gen_plan = False
    with st.spinner("Черчу карту подготовки к победе..."):
        context = f"Тема конкурса: {d['theme']}\n"
        if d['description']: context += f"Требования: {d['description']}\n"
        
        prompt = f"""
        Ты — наставник по проектным конкурсам. Твой ученик готовится к участию в конкурсе.
        {context}
        Идея, которую выбрал ученик: {d['selected']}

        Составь пошаговый план подготовки к конкурсу ИМЕННО ДЛЯ УЧАСТНИКА. Что ему нужно сделать, чтобы реализовать эту идею и достойно её защитить?
        Этапы: сбор материалов или исследование, создание работы или прототипа, подготовка к презентации, репетиция защиты. 
        Для каждого этапа укажи 2-3 конкретных действия и примерный срок в днях.
        Ответ оформи в виде структурированного списка.
        """
        d["plan"] = ask_ai(prompt)

if st.session_state.get("gen_risks"):
    st.session_state.gen_risks = False
    with st.spinner("Ищу подводные камни..."):
        prompt = f"""
        Ты — риск-менеджер. Твой ученик участвует в конкурсе с идеей: {d['selected']}
        Тема конкурса: {d['theme']}.

        Определи 5 ключевых рисков, с которыми может столкнуться ИМЕННО УЧАСТНИК при реализации этой идеи. Например: не хватит времени, подведет материал, сложно будет защитить проект, жесткие требования жюри.
        Для каждого риска укажи описание, вероятность низкая средняя или высокая, и способ снижения.
        Ответ верни в виде структурированного списка.
        """
        d["risks"] = ask_ai(prompt)

# Показываем красивые тексты на экране

st.title("Вектор 2035")
st.caption("От темы конкурса до готового плана подготовки")
st.write("«Вектор 2035» — ваш персональный ИИ-помощник для подготовки к любым конкурсам и проектам: от рисования до IT. Просто задайте тему, и нейросеть предложит креативные идеи, составит пошаговый план действий и предупредит о возможных рисках.")

# Вкладка 1 общая картина
if active == "1. План проекта":
    
    if not d['theme']:
        # Если проект пустой показываем простой текст и кнопочку
        st.write("")
        st.info("Проект ещё не создан. Начните с задания темы конкурса.")
        if st.button("Задать тему конкурса", type="primary"):
            go_to_tab("2. Отправка темы")
            
    else:
        # Если тема есть показываем всё по порядку
        st.header("Сводка по проекту")
        
        # Наша тема и параметры
        st.markdown(f"**Тема:** {d['theme']}")
        st.markdown(f"**Тип:** {d['category']}")
        if d['description']: st.markdown(f"**Описание:** {d['description']}")
        st.markdown(f"**Технологии:** {d['tech'] or '_без ограничений_'}")
        if d['audience']: st.markdown(f"**Аудитория:** {d['audience']}")
        
        st.divider()
        
        # Идеи которые придумал робот
        if d['ideas']:
            st.markdown("**Идеи:**")
            for i, idea in enumerate(d['ideas'], 1):
                clean_idea = re.sub(r'\*\*|__|#', '', idea).split("\n")[0]
                if clean_idea == d['selected']:
                    st.success(f"**Выбрано:** {clean_idea}")
                else:
                    st.markdown(f"{i}. {clean_idea}")
        else:
            st.info("Идеи ещё не сгенерированы.")
        
        st.divider()
        
        # План и риски прячем в выпадающие кнопочки чтобы страница не была огромной
        if d['plan']:
            with st.expander("Посмотреть план подготовки"):
                st.markdown(d['plan'])
        if d['risks']:
            with st.expander("Посмотреть риски реализации"):
                st.markdown(d['risks'])
        
        st.divider()
        
        # Кнопки внизу скачать файл или начать сначала
        col1, col2 = st.columns(2)
        with col1:
            docx_data = generate_docx(d)
            st.download_button(
                label="Скачать проект в Word (.docx)",
                data=docx_data,
                file_name=f"Проект_{d['theme'][:15].replace(' ', '_')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                type="primary",
                use_container_width=True
            )
        with col2:
            st.button("Начать проект заново", on_click=reset_project, use_container_width=True)

# Вкладка 2 придумываем тему
elif active == "2. Отправка темы":
    st.header("Параметры конкурса")
    
    d["theme"] = st.text_area("Тема конкурса обязательно", value=d["theme"])
    
    col1, col2 = st.columns(2)
    with col1:
        cats = ["Инновации / стартап", "Творческий", "Садоводство", "Моделирование", "Научный", "IT", "Спорт", "Другое"]
        d["category"] = st.selectbox("Тип конкурса", cats, index=cats.index(d["category"]))
    with col2:
        d["audience"] = st.text_input("Возрастная категория необязательно", value=d["audience"])
        
    d["description"] = st.text_area("Описание или требования необязательно", value=d["description"])
    d["tech"] = st.text_input("Технологии или материалы необязательно", value=d["tech"])
    
    st.button("Сохранить и перейти к идеям", type="primary", on_click=save_theme)
    
    if st.session_state.theme_error:
        st.error("Тема не может быть пустой!")

# Вкладка 3 выбираем идею
elif active == "3. Идеи":
    st.header("Идеи по вашей теме")
    if not d["theme"]:
        st.info("Сначала задайте тему конкурса чтобы мы могли сгенерировать идеи.")
        st.button("Перейти к вводу темы", on_click=go_to_tab, args=("2. Отправка темы",), type="primary")
    else:
        if st.button("Сгенерировать идеи заново"):
            st.session_state.gen_ideas = True
            st.rerun()
            
        for i, idea in enumerate(d["ideas"]):
            clean_idea = re.sub(r'\*\*|__|#', '', idea)
            with st.expander(f"Идея {i+1}"):
                st.write(clean_idea)
                st.button(f"Выбрать идею {i+1}", key=f"btn_{i}", on_click=select_idea, args=(clean_idea,))

# Вкладка 4 как готовиться план
elif active == "4. План подготовки":
    st.header("План подготовки")
    if not d["selected"]:
        st.info("Сначала выберите идею чтобы мы составили план подготовки.")
        st.button("Перейти к выбору идеи", on_click=go_to_tab, args=("3. Идеи",), type="primary")
    else:
        st.caption(f"Идея: {d['selected'][:150]}...")
        if st.button("Пересоздать план"):
            st.session_state.gen_plan = True
            st.rerun()
        st.markdown(d["plan"])

# Вкладка 5 чего бояться риски
elif active == "5. Риски":
    st.header("Риски реализации")
    if not d['selected']:
        st.info("Сначала выберите идею чтобы мы проанализировали риски.")
        st.button("Перейти к выбору идеи", on_click=go_to_tab, args=("3. Идеи",), type="primary")
    else:
        st.caption(f"Идея: {d['selected'][:150]}...")
        if st.button("Пересчитать риски"):
            st.session_state.gen_risks = True
            st.rerun()
        st.markdown(d["risks"])